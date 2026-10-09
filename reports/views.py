from __future__ import annotations
from django.shortcuts import render

# Create your views here.

from datetime import timedelta
from django.http import HttpResponse
from django.utils import timezone
from django.db.models import Sum, Count
import json
from core.authz import operator_required
from reports.services.services_reports import flights_count_by_period, revenue_by_period, revenue_by_route, top_flights_with_load_factor
from bookings.models import Booking

# za excel:
from io import BytesIO
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.styles import Font, Alignment


def _parse_date(request, key: str, default):
    raw = (request.GET.get(key) or "").strip()
    if not raw:
        return default
    try:
        return timezone.datetime.fromisoformat(raw).date()
    except Exception:
        return default


@operator_required
def operator_dashboard(request):
    today = timezone.localdate()
    default_from = today - timedelta(days=30)
    default_to = today

    from_date = _parse_date(request, "from", default_from)
    to_date = _parse_date(request, "to", default_to)

    group_by = (request.GET.get("group_by") or "day").strip().lower()
    if group_by not in ("day", "week", "month"):
        group_by = "day"

    flights_series = flights_count_by_period(from_date=from_date, to_date=to_date, group_by=group_by)
    revenue_series = revenue_by_period(from_date=from_date, to_date=to_date, group_by=group_by)
    route_table = revenue_by_route(from_date=from_date, to_date=to_date)
    top_flights = top_flights_with_load_factor(from_date=from_date, to_date=to_date, limit=20)

    totals = (
        Booking.objects.filter(
            status=Booking.Status.PAID,
            flight__scheduled_departure__date__range=(from_date, to_date),
        )
        .aggregate(
            revenue_sum=Sum("total_amount"),
            pax_sum=Sum("passengers_count"),
        )
    )
    revenue_sum = totals["revenue_sum"] or 0
    pax_sum = totals["pax_sum"] or 0

    flights_labels = [x["period"].strftime("%Y-%m-%d") if x.get("period") else "" for x in flights_series]
    flights_values = [int(x["cnt"]) for x in flights_series]

    revenue_labels = [x["period"].strftime("%Y-%m-%d") if x.get("period") else "" for x in revenue_series]
    revenue_values = [float(x.get("revenue") or 0) for x in revenue_series]
    revenue_pax_values = [int(x.get("pax") or 0) for x in revenue_series]

    top_routes = route_table[:10]
    routes_labels = [
        f'{r["flight__origin__iata_code"]}→{r["flight__destination__iata_code"]}'
        for r in top_routes
    ]
    routes_revenue = [float(r.get("revenue") or 0) for r in top_routes]
    routes_pax = [int(r.get("pax") or 0) for r in top_routes]

    top_flights_10 = top_flights[:10]
    top_flights_labels = [
        f"{f.flight_number} {f.origin}→{f.destination}"
        for f in top_flights_10
    ]
    top_flights_revenue = [float(f.revenue_paid or 0) for f in top_flights_10]
    top_flights_lf = [float(f.load_factor or 0) for f in top_flights_10]

    return render(
        request,
        "reports/operator_dashboard.html",
        {
            "from_date": from_date,
            "to_date": to_date,
            "group_by": group_by,

            "flights_series": flights_series,
            "revenue_series": revenue_series,
            "route_table": route_table,
            "top_flights": top_flights,

            "revenue_sum": revenue_sum,
            "pax_sum": pax_sum,

            "charts_json": json.dumps(
                {
                    "flights": {"labels": flights_labels, "values": flights_values},
                    "revenue": {
                        "labels": revenue_labels,
                        "values": revenue_values,
                        "pax": revenue_pax_values,
                    },
                    "routes": {
                        "labels": routes_labels,
                        "revenue": routes_revenue,
                        "pax": routes_pax,
                    },
                    "topFlights": {
                        "labels": top_flights_labels,
                        "revenue": top_flights_revenue,
                        "loadFactor": top_flights_lf,
                    },
                }
            ),
        },
    )



@operator_required
def export_bookings_csv(request):

    from datetime import timedelta

    today = timezone.localdate()
    default_from = today - timedelta(days=30)
    default_to = today

    from_date = _parse_date(request, "from", default_from)
    to_date = _parse_date(request, "to", default_to)

    qs = (
        Booking.objects.filter(
            status=Booking.Status.PAID,
            flight__scheduled_departure__date__range=(from_date, to_date),
        )
        .select_related("flight", "flight__origin", "flight__destination", "user")
        .order_by("-created_at")
    )

    # CSV
    resp = HttpResponse(content_type="text/csv")
    resp["Content-Disposition"] = f'attachment; filename="paid_bookings_{from_date}_{to_date}.csv"'

    import csv
    w = csv.writer(resp)
    w.writerow([
        "Booking Code",
        "Flight Number",
        "Airline",
        "Origin",
        "Destination",
        "Departure Date",
        "Departure Time",
        "Passengers",
        "Revenue (EUR)",
        "Booking Status",
        "Customer",
    ])
    for b in qs:
        flight = b.flight.flight_number
        airline = b.flight.airline.name
        origin = b.flight.origin.iata_code
        destination = b.flight.destination.iata_code

        dep_local = timezone.localtime(b.flight.scheduled_departure)
        dep_date = dep_local.date()
        dep_time = dep_local.time().strftime("%H:%M")

        w.writerow([
            b.booking_code,
            flight,
            airline,
            origin,
            destination,
            dep_date,
            dep_time,
            b.passengers_count,
            float(b.total_amount),
            b.status,
            b.user.username,
        ])
    return resp

@operator_required
def export_bookings_xlsx(request):

    from datetime import timedelta

    today = timezone.localdate()
    default_from = today - timedelta(days=30)
    default_to = today

    from_date = _parse_date(request, "from", default_from)
    to_date = _parse_date(request, "to", default_to)

    qs = (
        Booking.objects.filter(
            status=Booking.Status.PAID,
            flight__scheduled_departure__date__range=(from_date, to_date),
        )
        .select_related("flight", "flight__origin", "flight__destination", "user")
        .order_by("flight__scheduled_departure")
    )

    wb = Workbook()

    ws = wb.active
    ws.title = "Bookings"

    headers = [
        "Booking Code",
        "Flight",
        "Route",
        "Departure (Local)",
        "User",
        "Passengers",
        "Total Amount",
        "Currency",
        "Status",
    ]
    ws.append(headers)

    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for b in qs:
        dep_local = timezone.localtime(b.flight.scheduled_departure)
        route = f"{b.flight.origin.iata_code}→{b.flight.destination.iata_code}"
        ws.append([
            b.booking_code,
            b.flight.flight_number,
            route,
            dep_local.strftime("%Y-%m-%d %H:%M"),
            b.user.username,
            int(b.passengers_count or 0),
            float(b.total_amount or 0),
            b.currency,
            b.status,
        ])

    last_row = ws.max_row
    last_col = ws.max_column

    if last_row >= 2:
        table_ref = f"A1:{get_column_letter(last_col)}{last_row}"
        tab = Table(displayName="PaidBookings", ref=table_ref)
        style = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        tab.tableStyleInfo = style
        ws.add_table(tab)

    # Columns
    widths = [16, 10, 10, 18, 18, 10, 14, 10, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("Summary")
    totals = qs.aggregate(
        revenue_sum=Sum("total_amount"),
        pax_sum=Sum("passengers_count"),
        bookings_count=Count("id"),
    )
    revenue_sum = totals["revenue_sum"] or 0
    pax_sum = totals["pax_sum"] or 0
    bookings_count = totals["bookings_count"] or 0
    avg_rev = (revenue_sum / bookings_count) if bookings_count else 0

    ws2["A1"] = "Report"
    ws2["A1"].font = Font(bold=True, size=14)

    ws2["A3"] = "From"
    ws2["B3"] = str(from_date)
    ws2["A4"] = "To"
    ws2["B4"] = str(to_date)

    ws2["A6"] = "Total revenue"
    ws2["B6"] = float(revenue_sum)
    ws2["A7"] = "Total passengers"
    ws2["B7"] = int(pax_sum)
    ws2["A8"] = "Paid bookings"
    ws2["B8"] = int(bookings_count)
    ws2["A9"] = "Avg revenue/booking"
    ws2["B9"] = float(avg_rev)

    for r in range(6, 10):
        ws2[f"A{r}"].font = Font(bold=True)

    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 18

    ws3 = wb.create_sheet("Charts")

    group_by = (request.GET.get("group_by") or "day").strip().lower()
    if group_by not in ("day", "week", "month"):
        group_by = "day"

    flights_series = flights_count_by_period(from_date=from_date, to_date=to_date, group_by=group_by)
    revenue_series = revenue_by_period(from_date=from_date, to_date=to_date, group_by=group_by)

    ws3.append(["Period", "Flights", "Revenue", "Passengers"])
    ws3["A1"].font = Font(bold=True)
    ws3["B1"].font = Font(bold=True)
    ws3["C1"].font = Font(bold=True)
    ws3["D1"].font = Font(bold=True)

    f_map = {x["period"]: int(x["cnt"]) for x in flights_series}
    for x in revenue_series:
        period = x["period"]
        ws3.append([
            period.strftime("%Y-%m-%d") if period else "",
            f_map.get(period, 0),
            float(x.get("revenue") or 0),
            int(x.get("pax") or 0),
        ])

    if ws3.max_row >= 3:
        line = LineChart()
        line.title = "Revenue by period"
        line.y_axis.title = "Revenue"
        line.x_axis.title = "Period"

        data = Reference(ws3, min_col=3, min_row=1, max_row=ws3.max_row)
        cats = Reference(ws3, min_col=1, min_row=2, max_row=ws3.max_row)
        line.add_data(data, titles_from_data=True)
        line.set_categories(cats)
        ws3.add_chart(line, "F2")

        bar = BarChart()
        bar.title = "Flights count by period"
        bar.y_axis.title = "Flights"
        bar.x_axis.title = "Period"
        data2 = Reference(ws3, min_col=2, min_row=1, max_row=ws3.max_row)
        bar.add_data(data2, titles_from_data=True)
        bar.set_categories(cats)
        ws3.add_chart(bar, "F18")

    ws3.column_dimensions["A"].width = 14
    ws3.column_dimensions["B"].width = 10
    ws3.column_dimensions["C"].width = 14
    ws3.column_dimensions["D"].width = 12

    out = BytesIO()
    wb.save(out)
    out.seek(0)

    resp = HttpResponse(
        out.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f'attachment; filename="paid_bookings_{from_date}_{to_date}.xlsx"'
    return resp